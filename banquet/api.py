from django.http import HttpResponse
from django.utils import timezone
from ninja import NinjaAPI, Query
from datetime import timedelta
import csv
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

from .models import (
    BanquetBooking, BookingStatus, Employee, Role,
    MenuItem, MenuConfirmation, MenuConfirmationItem, StuckLevel, AuditLog
)
from .schemas import (
    BookingSchema, BookingDetailSchema, BookingCreateSchema, BookingUpdateSchema,
    MenuConfirmSchema, StateOperationSchema, EmployeeSchema, MenuItemSchema,
    StuckBookingSchema, ExportRequestSchema, AuditLogSchema, MenuConfirmationSchema
)
from .services import (
    submit_booking, confirm_menu, receive_by_kitchen,
    start_service, complete_booking, cancel_booking,
    get_stuck_bookings, StateTransitionError
)
from .seed import create_seed_data


api = NinjaAPI(
    title='宴会管理系统 API',
    description='酒店宴会部宴会预订与菜单确认系统',
    version='1.0.0'
)


def _get_client_ip(request):
    x_forwarded_for = request.META.get('HTTP_X_FORWARDED_FOR')
    if x_forwarded_for:
        return x_forwarded_for.split(',')[0]
    return request.META.get('REMOTE_ADDR')


def _serialize_booking(booking, include_stuck_info=False):
    data = {
        'id': booking.id,
        'booking_no': booking.booking_no,
        'customer_name': booking.customer_name,
        'customer_phone': booking.customer_phone,
        'banquet_type': booking.banquet_type,
        'banquet_date': booking.banquet_date,
        'start_time': booking.start_time,
        'end_time': booking.end_time,
        'venue': booking.venue,
        'expected_guests': booking.expected_guests,
        'table_count': booking.table_count,
        'budget_per_table': float(booking.budget_per_table),
        'status': booking.status,
        'status_display': booking.get_status_display(),
        'sales_person_id': booking.sales_person_id,
        'sales_person_name': booking.sales_person.name,
        'floor_supervisor_id': booking.floor_supervisor_id,
        'floor_supervisor_name': booking.floor_supervisor.name if booking.floor_supervisor else None,
        'kitchen_coordinator_id': booking.kitchen_coordinator_id,
        'kitchen_coordinator_name': booking.kitchen_coordinator.name if booking.kitchen_coordinator else None,
        'remarks': booking.remarks,
        'submitted_at': booking.submitted_at,
        'created_at': booking.created_at,
        'updated_at': booking.updated_at,
    }
    if include_stuck_info:
        stuck_level = booking.get_stuck_level()
        data.update({
            'stuck_level': stuck_level,
            'stuck_level_display': StuckLevel(stuck_level).label,
            'stuck_deadline': booking.get_stuck_deadline(),
        })
    return data


def _serialize_menu_confirmation(mc):
    items = MenuConfirmationItem.objects.filter(menu_confirmation=mc).select_related('menu_item')
    return {
        'id': mc.id,
        'booking_id': mc.booking_id,
        'total_amount': float(mc.total_amount),
        'special_requirements': mc.special_requirements,
        'wine_arrangement': mc.wine_arrangement,
        'table_layout': mc.table_layout,
        'confirmed_by_id': mc.confirmed_by_id,
        'confirmed_by_name': mc.confirmed_by.name,
        'confirmed_at': mc.confirmed_at,
        'customer_signed': mc.customer_signed,
        'customer_signature': mc.customer_signature,
        'remarks': mc.remarks,
        'items': [
            {
                'id': item.id,
                'menu_item_id': item.menu_item_id,
                'menu_item_name': item.menu_item.name,
                'category': item.menu_item.category,
                'quantity': item.quantity,
                'unit_price': float(item.unit_price),
                'subtotal': float(item.subtotal),
                'remarks': item.remarks,
            }
            for item in items
        ],
        'created_at': mc.created_at,
        'updated_at': mc.updated_at,
    }


def _serialize_audit_log(log):
    return {
        'id': log.id,
        'booking_id': log.booking_id,
        'action': log.action,
        'from_status': log.from_status,
        'from_status_display': BookingStatus(log.from_status).label if log.from_status else None,
        'to_status': log.to_status,
        'to_status_display': BookingStatus(log.to_status).label if log.to_status else None,
        'operator_id': log.operator_id,
        'operator_name': log.operator.name,
        'operator_role': log.operator_role,
        'operator_role_display': Role(log.operator_role).label,
        'timestamp': log.timestamp,
        'remarks': log.remarks,
    }


@api.get('/employees', response=list[EmployeeSchema], tags=['基础数据'])
def list_employees(request, role: str | None = None):
    queryset = Employee.objects.filter(is_active=True)
    if role:
        queryset = queryset.filter(role=role)
    return [
        {
            'id': e.id,
            'name': e.name,
            'role': e.role,
            'role_display': e.get_role_display(),
            'phone': e.phone,
        }
        for e in queryset
    ]


@api.get('/menu-items', response=list[MenuItemSchema], tags=['基础数据'])
def list_menu_items(request, category: str | None = None):
    queryset = MenuItem.objects.filter(is_active=True)
    if category:
        queryset = queryset.filter(category=category)
    return [
        {
            'id': m.id,
            'name': m.name,
            'category': m.category,
            'price': float(m.price),
            'description': m.description,
        }
        for m in queryset
    ]


@api.get('/bookings', response=list[BookingSchema], tags=['宴会预订'])
def list_bookings(
    request,
    status: str | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
    sales_person_id: int | None = None,
    include_stuck_info: bool = True,
):
    queryset = BanquetBooking.objects.select_related(
        'sales_person', 'floor_supervisor', 'kitchen_coordinator'
    )
    if status:
        queryset = queryset.filter(status=status)
    if start_date:
        queryset = queryset.filter(banquet_date__gte=start_date)
    if end_date:
        queryset = queryset.filter(banquet_date__lte=end_date)
    if sales_person_id:
        queryset = queryset.filter(sales_person_id=sales_person_id)

    return [_serialize_booking(b, include_stuck_info) for b in queryset]


@api.get('/bookings/{booking_id}', response=BookingDetailSchema, tags=['宴会预订'])
def get_booking(request, booking_id: int):
    booking = BanquetBooking.objects.select_related(
        'sales_person', 'floor_supervisor', 'kitchen_coordinator'
    ).get(id=booking_id)

    result = _serialize_booking(booking, include_stuck_info=True)

    try:
        mc = MenuConfirmation.objects.select_related('confirmed_by').get(booking=booking)
        result['menu_confirmation'] = _serialize_menu_confirmation(mc)
    except MenuConfirmation.DoesNotExist:
        result['menu_confirmation'] = None

    audit_logs = AuditLog.objects.filter(booking=booking).select_related('operator').order_by('-timestamp')
    result['audit_logs'] = [_serialize_audit_log(log) for log in audit_logs]

    return result


@api.post('/bookings', response=BookingSchema, tags=['宴会预订'])
def create_booking(request, payload: BookingCreateSchema):
    latest = BanquetBooking.objects.order_by('-id').first()
    seq = 1
    if latest:
        seq = int(latest.booking_no.split('-')[-1]) + 1
    booking_no = f'BY{timezone.now().strftime("%Y%m%d")}-{seq:04d}'

    booking = BanquetBooking.objects.create(
        booking_no=booking_no,
        **payload.model_dump()
    )
    return _serialize_booking(booking)


@api.put('/bookings/{booking_id}', response=BookingSchema, tags=['宴会预订'])
def update_booking(request, booking_id: int, payload: BookingUpdateSchema):
    booking = BanquetBooking.objects.get(id=booking_id)
    if booking.status != BookingStatus.DRAFT:
        return api.create_response(request, {'detail': '只能编辑草稿状态的预订'}, status=400)

    for key, value in payload.model_dump(exclude_unset=True).items():
        setattr(booking, key, value)
    booking.save()
    return _serialize_booking(booking)


@api.post('/bookings/{booking_id}/submit', tags=['状态流转'])
def api_submit_booking(request, booking_id: int, payload: StateOperationSchema):
    try:
        booking = submit_booking(
            booking_id=booking_id,
            operator_id=payload.operator_id,
            remarks=payload.remarks,
            ip_address=_get_client_ip(request)
        )
        return {'success': True, 'booking': _serialize_booking(booking)}
    except StateTransitionError as e:
        return api.create_response(request, {'success': False, 'detail': str(e)}, status=400)


@api.post('/bookings/{booking_id}/confirm-menu', tags=['状态流转'])
def api_confirm_menu(request, booking_id: int, payload: MenuConfirmSchema):
    operator_id = int(request.headers.get('X-Operator-Id', 0))
    if not operator_id:
        return api.create_response(request, {'detail': '请在请求头 X-Operator-Id 中指定操作人ID'}, status=400)

    try:
        booking, mc = confirm_menu(
            booking_id=booking_id,
            operator_id=operator_id,
            menu_data=payload.model_dump(),
            remarks=payload.remarks,
            ip_address=_get_client_ip(request)
        )
        return {
            'success': True,
            'booking': _serialize_booking(booking),
            'menu_confirmation': _serialize_menu_confirmation(mc),
        }
    except StateTransitionError as e:
        return api.create_response(request, {'success': False, 'detail': str(e)}, status=400)


@api.post('/bookings/{booking_id}/kitchen-receive', tags=['状态流转'])
def api_kitchen_receive(request, booking_id: int, payload: StateOperationSchema):
    try:
        booking = receive_by_kitchen(
            booking_id=booking_id,
            operator_id=payload.operator_id,
            remarks=payload.remarks,
            ip_address=_get_client_ip(request)
        )
        return {'success': True, 'booking': _serialize_booking(booking)}
    except StateTransitionError as e:
        return api.create_response(request, {'success': False, 'detail': str(e)}, status=400)


@api.post('/bookings/{booking_id}/start', tags=['状态流转'])
def api_start_service(request, booking_id: int, payload: StateOperationSchema):
    try:
        booking = start_service(
            booking_id=booking_id,
            operator_id=payload.operator_id,
            remarks=payload.remarks,
            ip_address=_get_client_ip(request)
        )
        return {'success': True, 'booking': _serialize_booking(booking)}
    except StateTransitionError as e:
        return api.create_response(request, {'success': False, 'detail': str(e)}, status=400)


@api.post('/bookings/{booking_id}/complete', tags=['状态流转'])
def api_complete_booking(request, booking_id: int, payload: StateOperationSchema):
    try:
        booking = complete_booking(
            booking_id=booking_id,
            operator_id=payload.operator_id,
            remarks=payload.remarks,
            ip_address=_get_client_ip(request)
        )
        return {'success': True, 'booking': _serialize_booking(booking)}
    except StateTransitionError as e:
        return api.create_response(request, {'success': False, 'detail': str(e)}, status=400)


@api.post('/bookings/{booking_id}/cancel', tags=['状态流转'])
def api_cancel_booking(request, booking_id: int, payload: StateOperationSchema):
    try:
        booking = cancel_booking(
            booking_id=booking_id,
            operator_id=payload.operator_id,
            remarks=payload.remarks,
            ip_address=_get_client_ip(request)
        )
        return {'success': True, 'booking': _serialize_booking(booking)}
    except StateTransitionError as e:
        return api.create_response(request, {'success': False, 'detail': str(e)}, status=400)


@api.get('/bookings/{booking_id}/audit-logs', response=list[AuditLogSchema], tags=['审计日志'])
def list_audit_logs(request, booking_id: int):
    logs = AuditLog.objects.filter(booking_id=booking_id).select_related('operator').order_by('-timestamp')
    return [_serialize_audit_log(log) for log in logs]


@api.get('/stuck-bookings', response=list[StuckBookingSchema], tags=['卡住预警'])
def api_stuck_bookings(request):
    stuck_data = get_stuck_bookings()
    return [
        {
            'booking': _serialize_booking(item['booking'], include_stuck_info=True),
            'stuck_level': item['stuck_level'],
            'stuck_level_display': StuckLevel(item['stuck_level']).label,
            'deadline': item['deadline'],
            'overdue_hours': item['overdue_hours'],
        }
        for item in stuck_data
    ]


@api.get('/bookings/{booking_id}/menu-confirmation', response=MenuConfirmationSchema, tags=['菜单确认回看'])
def get_menu_confirmation(request, booking_id: int):
    mc = MenuConfirmation.objects.select_related('confirmed_by').get(booking_id=booking_id)
    return _serialize_menu_confirmation(mc)


@api.post('/export/bookings', tags=['导出任务'])
def export_bookings(request, payload: ExportRequestSchema):
    queryset = BanquetBooking.objects.select_related(
        'sales_person', 'floor_supervisor', 'kitchen_coordinator'
    )

    if payload.start_date:
        queryset = queryset.filter(banquet_date__gte=payload.start_date)
    if payload.end_date:
        queryset = queryset.filter(banquet_date__lte=payload.end_date)
    if payload.status:
        queryset = queryset.filter(status=payload.status)

    bookings = list(queryset)
    if payload.include_stuck_only:
        bookings = [b for b in bookings if b.get_stuck_level() != StuckLevel.NORMAL]

    if payload.format == 'csv':
        return _export_csv(bookings)
    else:
        return _export_xlsx(bookings)


def _export_csv(bookings):
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="banquet_bookings_{timezone.now().strftime("%Y%m%d_%H%M%S")}.csv"'
    response.write('\ufeff')

    writer = csv.writer(response)
    headers = [
        '预订单号', '客户姓名', '联系电话', '宴会类型', '宴会日期',
        '开始时间', '结束时间', '场地', '预计人数', '桌数',
        '每桌预算', '状态', '卡住状态', '销售人员', '厅面主管',
        '后厨统筹', '提交时间', '创建时间'
    ]
    writer.writerow(headers)

    for b in bookings:
        stuck_level = b.get_stuck_level()
        writer.writerow([
            b.booking_no, b.customer_name, b.customer_phone, b.banquet_type, b.banquet_date,
            b.start_time, b.end_time, b.venue, b.expected_guests, b.table_count,
            float(b.budget_per_table), b.get_status_display(), StuckLevel(stuck_level).label,
            b.sales_person.name,
            b.floor_supervisor.name if b.floor_supervisor else '',
            b.kitchen_coordinator.name if b.kitchen_coordinator else '',
            b.submitted_at if b.submitted_at else '',
            b.created_at.strftime('%Y-%m-%d %H:%M:%S'),
        ])

    return response


def _export_xlsx(bookings):
    wb = Workbook()
    ws = wb.active
    ws.title = '宴会预订列表'

    headers = [
        '预订单号', '客户姓名', '联系电话', '宴会类型', '宴会日期',
        '开始时间', '结束时间', '场地', '预计人数', '桌数',
        '每桌预算', '状态', '卡住状态', '销售人员', '厅面主管',
        '后厨统筹', '提交时间', '创建时间'
    ]

    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row, b in enumerate(bookings, 2):
        stuck_level = b.get_stuck_level()
        stuck_display = StuckLevel(stuck_level).label
        status_display = b.get_status_display()

        row_data = [
            b.booking_no, b.customer_name, b.customer_phone, b.banquet_type, str(b.banquet_date),
            str(b.start_time), str(b.end_time), b.venue, b.expected_guests, b.table_count,
            float(b.budget_per_table), status_display, stuck_display,
            b.sales_person.name,
            b.floor_supervisor.name if b.floor_supervisor else '',
            b.kitchen_coordinator.name if b.kitchen_coordinator else '',
            str(b.submitted_at) if b.submitted_at else '',
            b.created_at.strftime('%Y-%m-%d %H:%M:%S'),
        ]

        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row, column=col, value=value)
            if stuck_level == StuckLevel.STUCK:
                cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
            elif stuck_level == StuckLevel.WARNING:
                cell.fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column].width = min(max_length + 2, 30)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="banquet_bookings_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
    return response


@api.post('/seed', tags=['系统'])
def api_create_seed_data(request, reset: bool = False):
    result = create_seed_data(reset=reset)
    return result


@api.get('/status-flow', tags=['系统'])
def get_status_flow(request):
    return {
        'statuses': [
            {'value': s.value, 'label': s.label} for s in BookingStatus
        ],
        'roles': [
            {'value': r.value, 'label': r.label} for r in Role
        ],
        'transitions': [
            {'from': 'draft', 'to': 'submitted', 'action': '提交预订', 'role': 'sales', 'description': '宴会销售提交预订给厅面主管'},
            {'from': 'submitted', 'to': 'menu_confirmed', 'action': '确认菜单', 'role': 'floor_supervisor', 'description': '厅面主管确认菜单并安排'},
            {'from': 'menu_confirmed', 'to': 'kitchen_received', 'action': '后厨接收', 'role': 'kitchen_coordinator', 'description': '后厨统筹接收菜单安排'},
            {'from': 'kitchen_received', 'to': 'in_progress', 'action': '开始服务', 'role': 'any', 'description': '宴会当天开始服务'},
            {'from': 'in_progress', 'to': 'completed', 'action': '完成', 'role': 'any', 'description': '宴会结束'},
            {'from': '*', 'to': 'cancelled', 'action': '取消', 'role': 'any', 'description': '取消预订'},
        ],
        'timeout_config': {
            'menu_confirmation_timeout_hours': 24,
            'kitchen_receive_timeout_hours': 12,
            'description': '提交预订后24小时内需确认菜单，菜单确认后12小时内需后厨接收，超时标记为卡住'
        }
    }
