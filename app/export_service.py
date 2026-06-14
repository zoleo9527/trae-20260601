from sqlalchemy.orm import Session
from sqlalchemy import and_
from app.models import ExportTask, ActivityMaterial, StoreFeedback, User, Store
from app.schemas import ExportTaskCreate
from datetime import datetime
from typing import Optional, List
import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
import os


class ExportService:
    EXPORT_DIR = "exports"

    @staticmethod
    def ensure_export_dir():
        if not os.path.exists(ExportService.EXPORT_DIR):
            os.makedirs(ExportService.EXPORT_DIR)

    @staticmethod
    def create_export_task(db: Session, task: ExportTaskCreate, user_id: int) -> ExportTask:
        ExportService.ensure_export_dir()
        db_task = ExportTask(**task.model_dump(), created_by_id=user_id, status="pending")
        db.add(db_task)
        db.commit()
        db.refresh(db_task)
        return db_task

    @staticmethod
    def get_export_task(db: Session, task_id: int) -> Optional[ExportTask]:
        return db.query(ExportTask).filter(ExportTask.id == task_id).first()

    @staticmethod
    def get_export_tasks(
        db: Session, skip: int = 0, limit: int = 100, created_by_id: Optional[int] = None
    ) -> List[ExportTask]:
        query = db.query(ExportTask)
        if created_by_id:
            query = query.filter(ExportTask.created_by_id == created_by_id)
        return query.offset(skip).limit(limit).all()

    @staticmethod
    def execute_export_task(db: Session, task_id: int) -> Optional[ExportTask]:
        task = ExportService.get_export_task(db, task_id)
        if not task or task.status != "pending":
            return task

        try:
            task.status = "processing"
            db.commit()

            if task.task_type == "materials":
                file_path = ExportService.export_materials(db, task.parameters)
            elif task.task_type == "feedbacks":
                file_path = ExportService.export_feedbacks(db, task.parameters)
            elif task.task_type == "stuck_items":
                file_path = ExportService.export_stuck_items(db, task.parameters)
            else:
                raise ValueError(f"Unknown task type: {task.task_type}")

            task.status = "completed"
            task.file_path = file_path
            task.completed_at = datetime.now()
            db.commit()
            db.refresh(task)

        except Exception as e:
            task.status = "failed"
            task.error_message = str(e)
            db.commit()
            db.refresh(task)

        return task

    @staticmethod
    def export_materials(db: Session, parameters: Optional[str]) -> str:
        ExportService.ensure_export_dir()
        wb = Workbook()
        ws = wb.active
        ws.title = "活动物料"

        headers = [
            "ID",
            "物料名称",
            "物料编号",
            "门店",
            "数量",
            "状态",
            "当前处理人",
            "处理人角色",
            "发放时间",
            "接收时间",
            "预计完成时间",
            "实际完成时间",
            "卡住原因",
            "创建时间",
            "更新时间",
        ]
        ws.append(headers)

        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        filters = {}
        if parameters:
            try:
                filters = json.loads(parameters)
            except:
                pass

        query = db.query(ActivityMaterial)
        if "store_id" in filters:
            query = query.filter(ActivityMaterial.store_id == filters["store_id"])
        if "status" in filters:
            query = query.filter(ActivityMaterial.status == filters["status"])

        materials = query.all()

        for material in materials:
            store = db.query(Store).filter(Store.id == material.store_id).first()
            handler = None
            handler_role = ""
            if material.current_handler_id:
                handler = db.query(User).filter(User.id == material.current_handler_id).first()
                if handler:
                    handler_role = handler.role.value

            row = [
                material.id,
                material.name,
                material.code,
                store.name if store else "",
                material.quantity,
                material.status.value,
                handler.real_name if handler else "",
                handler_role,
                material.distributed_at.isoformat() if material.distributed_at else "",
                material.received_at.isoformat() if material.received_at else "",
                material.expected_complete_date.isoformat()
                if material.expected_complete_date
                else "",
                material.actual_complete_date.isoformat() if material.actual_complete_date else "",
                material.stuck_reason or "",
                material.created_at.isoformat() if material.created_at else "",
                material.updated_at.isoformat() if material.updated_at else "",
            ]
            ws.append(row)

        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2) * 1.2
            ws.column_dimensions[column_letter].width = adjusted_width

        filename = f"materials_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        file_path = os.path.join(ExportService.EXPORT_DIR, filename)
        wb.save(file_path)

        return file_path

    @staticmethod
    def export_feedbacks(db: Session, parameters: Optional[str]) -> str:
        ExportService.ensure_export_dir()
        wb = Workbook()
        ws = wb.active
        ws.title = "门店反馈"

        headers = [
            "ID",
            "标题",
            "内容",
            "关联物料",
            "门店",
            "反馈类型",
            "优先级",
            "状态",
            "当前处理人",
            "处理人角色",
            "解决方案",
            "拒绝原因",
            "卡住原因",
            "创建时间",
            "更新时间",
            "解决时间",
        ]
        ws.append(headers)

        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        filters = {}
        if parameters:
            try:
                filters = json.loads(parameters)
            except:
                pass

        query = db.query(StoreFeedback)
        if "store_id" in filters:
            query = query.filter(StoreFeedback.store_id == filters["store_id"])
        if "material_id" in filters:
            query = query.filter(StoreFeedback.material_id == filters["material_id"])
        if "status" in filters:
            query = query.filter(StoreFeedback.status == filters["status"])

        feedbacks = query.all()

        for feedback in feedbacks:
            material = (
                db.query(ActivityMaterial)
                .filter(ActivityMaterial.id == feedback.material_id)
                .first()
            )
            store = db.query(Store).filter(Store.id == feedback.store_id).first()
            handler = None
            handler_role = ""
            if feedback.current_handler_id:
                handler = db.query(User).filter(User.id == feedback.current_handler_id).first()
                if handler:
                    handler_role = handler.role.value

            row = [
                feedback.id,
                feedback.title,
                feedback.content,
                material.name if material else "",
                store.name if store else "",
                feedback.feedback_type or "",
                feedback.priority,
                feedback.status.value,
                handler.real_name if handler else "",
                handler_role,
                feedback.resolution or "",
                feedback.rejected_reason or "",
                feedback.stuck_reason or "",
                feedback.created_at.isoformat() if feedback.created_at else "",
                feedback.updated_at.isoformat() if feedback.updated_at else "",
                feedback.resolved_at.isoformat() if feedback.resolved_at else "",
            ]
            ws.append(row)

        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2) * 1.2
            ws.column_dimensions[column_letter].width = adjusted_width

        filename = f"feedbacks_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        file_path = os.path.join(ExportService.EXPORT_DIR, filename)
        wb.save(file_path)

        return file_path

    @staticmethod
    def export_stuck_items(db: Session, parameters: Optional[str]) -> str:
        ExportService.ensure_export_dir()
        wb = Workbook()
        ws = wb.active
        ws.title = "卡住项目"

        headers = [
            "类型",
            "ID",
            "名称",
            "门店",
            "当前状态",
            "当前处理人",
            "处理人角色",
            "卡住原因",
            "卡住天数",
            "创建时间",
            "更新时间",
        ]
        ws.append(headers)

        header_fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        filters = {}
        if parameters:
            try:
                filters = json.loads(parameters)
            except:
                pass

        days = filters.get("days", 3)

        from app.services import MaterialService, FeedbackService

        stuck_materials = MaterialService.get_stuck_materials(db, days)
        for material in stuck_materials:
            store = db.query(Store).filter(Store.id == material.store_id).first()
            handler = None
            handler_role = ""
            if material.current_handler_id:
                handler = db.query(User).filter(User.id == material.current_handler_id).first()
                if handler:
                    handler_role = handler.role.value

            days_stuck = 0
            if material.updated_at:
                days_stuck = (datetime.now() - material.updated_at).days

            row = [
                "活动物料",
                material.id,
                material.name,
                store.name if store else "",
                material.status.value,
                handler.real_name if handler else "",
                handler_role,
                material.stuck_reason or "",
                days_stuck,
                material.created_at.isoformat() if material.created_at else "",
                material.updated_at.isoformat() if material.updated_at else "",
            ]
            ws.append(row)

        stuck_feedbacks = FeedbackService.get_stuck_feedbacks(db, days)
        for feedback in stuck_feedbacks:
            material = (
                db.query(ActivityMaterial)
                .filter(ActivityMaterial.id == feedback.material_id)
                .first()
            )
            store = db.query(Store).filter(Store.id == feedback.store_id).first()
            handler = None
            handler_role = ""
            if feedback.current_handler_id:
                handler = db.query(User).filter(User.id == feedback.current_handler_id).first()
                if handler:
                    handler_role = handler.role.value

            days_stuck = 0
            if feedback.updated_at:
                days_stuck = (datetime.now() - feedback.updated_at).days

            row = [
                "门店反馈",
                feedback.id,
                feedback.title,
                store.name if store else "",
                feedback.status.value,
                handler.real_name if handler else "",
                handler_role,
                feedback.stuck_reason or "",
                days_stuck,
                feedback.created_at.isoformat() if feedback.created_at else "",
                feedback.updated_at.isoformat() if feedback.updated_at else "",
            ]
            ws.append(row)

        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2) * 1.2
            ws.column_dimensions[column_letter].width = adjusted_width

        filename = f"stuck_items_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        file_path = os.path.join(ExportService.EXPORT_DIR, filename)
        wb.save(file_path)

        return file_path