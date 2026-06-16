"""
母婴零售店 - Django Ninja API接口
包含仪表板、退换货、客户回访、审计日志、导出等核心功能
"""
from django.http import HttpResponse
from django.utils import timezone
from django.db.models import Count, Q
from datetime import datetime, timedelta
from ninja import Router, Query, Field
from typing import Optional, List
from pydantic import BaseModel, ConfigDict
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from io import BytesIO

from .models import (
    Employee, Member, Product, ProductBatch, ReturnExchange, VisitRecord,
    AuditLog, Notification, Attachment, PromotionCoupon
)

router = Router()


class MemberSchema(BaseModel):
    model_config = ConfigDict(from_attributes=True)

    member_id: str
    name: str
    phone: str
    member_level: str


class EmployeeSchema(BaseModel):
    model_config = ConfigDict(from_attributes=True)

    id: int
    employee_id: str
    role: str
    user__username: str
    user__first_name: str
    user__last_name: str


class DashboardReturnItem(BaseModel):
    model_config = ConfigDict(from_attributes=True)

    id: int
    return_number: str
    member_name: str
    type: str
    status: str
    reason_category: str
    assigned_to_name: Optional[str] = None
    current_handler_role: str
    stuck_info: Optional[dict] = None
    created_at: datetime
    days_pending: int


class DashboardVisitItem(BaseModel):
    model_config = ConfigDict(from_attributes=True)

    id: int
    visit_number: str
    member_name: str
    type: str
    priority: str
    status: str
    purpose: str
    assigned_to_name: Optional[str] = None
    scheduled_date: datetime
    scheduled_time: Optional[datetime] = None
    stuck_info: Optional[dict] = None
    days_overdue: Optional[int] = None


class DashboardResponse(BaseModel):
    today_returns_pending: int
    today_visits_pending: int
    stuck_returns: int
    stuck_visits: int
    returns_needing_attention: List[DashboardReturnItem]
    visits_needing_attention: List[DashboardVisitItem]
    my_pending_returns: List[DashboardReturnItem]
    my_pending_visits: List[DashboardVisitItem]
    returns_by_role: dict
    visits_by_role: dict
    statistics: dict


@router.get("/dashboard", response=DashboardResponse)
def get_dashboard(request):
    """
    仪表板接口 - 展示今天应该先处理什么
    回答三个核心问题：
    1. 谁在处理（店员、店长、采购）
    2. 退换货卡在哪里
    3. 客户回访为什么还没完成
    """
    today = timezone.now().date()

    returns = ReturnExchange.objects.select_related(
        'member', 'assigned_to', 'assigned_to__user'
    ).exclude(status__in=['completed', 'rejected'])

    visits = VisitRecord.objects.select_related(
        'member', 'assigned_to', 'assigned_to__user'
    ).filter(
        Q(scheduled_date=today) | Q(scheduled_date__lt=today),
        status__in=['pending', 'in_progress']
    )

    all_returns_list = []
    stuck_returns_count = 0
    returns_by_role_data = {
        'clerk': [],
        'manager': [],
        'purchaser': [],
    }

    for ret in returns:
        stuck_info = ret.get_stuck_info()
        item = {
            'id': ret.id,
            'return_number': ret.return_number,
            'member_name': ret.member.name,
            'type': ret.get_type_display(),
            'status': ret.get_status_display(),
            'reason_category': ret.get_reason_category_display(),
            'assigned_to_name': ret.assigned_to.user.get_full_name() if ret.assigned_to else None,
            'current_handler_role': ret.get_current_expected_handler(),
            'stuck_info': stuck_info,
            'created_at': ret.created_at,
            'days_pending': (timezone.now() - ret.created_at).days,
        }
        all_returns_list.append(item)

        if stuck_info and stuck_info.get('is_stuck'):
            stuck_returns_count += 1

        handler_role = ret.get_current_expected_handler()
        if handler_role in returns_by_role_data:
            returns_by_role_data[handler_role].append(item)

    all_visits_list = []
    stuck_visits_count = 0
    visits_by_role_data = {
        'clerk': [],
        'manager': [],
        'purchaser': [],
    }

    for visit in visits:
        stuck_info = visit.get_stuck_info()
        days_overdue = None
        if stuck_info and stuck_info.get('is_stuck'):
            days_overdue = stuck_info.get('days_overdue', 0)

        item = {
            'id': visit.id,
            'visit_number': visit.visit_number,
            'member_name': visit.member.name,
            'type': visit.get_type_display(),
            'priority': visit.get_priority_display(),
            'status': visit.get_status_display(),
            'purpose': visit.purpose,
            'assigned_to_name': visit.assigned_to.user.get_full_name() if visit.assigned_to else None,
            'scheduled_date': visit.scheduled_date,
            'scheduled_time': visit.scheduled_time,
            'stuck_info': stuck_info,
            'days_overdue': days_overdue,
        }
        all_visits_list.append(item)

        if stuck_info and stuck_info.get('is_stuck'):
            stuck_visits_count += 1

        if visit.assigned_to:
            role = visit.assigned_to.role
            if role in visits_by_role_data:
                visits_by_role_data[role].append(item)

    returns_needing_attention = sorted(
        [r for r in all_returns_list if r['stuck_info'] and r['stuck_info'].get('is_stuck')],
        key=lambda x: x['days_pending'],
        reverse=True
    )[:10]

    visits_needing_attention = sorted(
        [v for v in all_visits_list if v['stuck_info'] and v['stuck_info'].get('is_stuck')],
        key=lambda x: x['days_overdue'] or 0,
        reverse=True
    )[:10]

    my_returns = []
    my_visits = []
    current_employee = None
    if hasattr(request, 'user') and request.user.is_authenticated:
        try:
            current_employee = request.user.employee_profile
            my_returns = [
                item for item in all_returns_list
                if item['assigned_to_name'] == current_employee.user.get_full_name()
            ][:5]
            my_visits = [
                item for item in all_visits_list
                if item['assigned_to_name'] == current_employee.user.get_full_name()
            ][:5]
        except Employee.DoesNotExist:
            pass

    stats = {
        'total_returns_pending': returns.count(),
        'total_visits_pending': visits.count(),
        'today_returns_created': ReturnExchange.objects.filter(
            created_at__date=today
        ).count(),
        'today_returns_completed': ReturnExchange.objects.filter(
            completed_at__date=today, status='completed'
        ).count(),
        'today_visits_scheduled': VisitRecord.objects.filter(
            scheduled_date=today
        ).count(),
        'today_visits_completed': VisitRecord.objects.filter(
            completed_date=today, status='completed'
        ).count(),
    }

    return {
        'today_returns_pending': returns.count(),
        'today_visits_pending': visits.count(),
        'stuck_returns': stuck_returns_count,
        'stuck_visits': stuck_visits_count,
        'returns_needing_attention': returns_needing_attention,
        'visits_needing_attention': visits_needing_attention,
        'my_pending_returns': my_returns,
        'my_pending_visits': my_visits,
        'returns_by_role': {
            'clerk': {
                'count': len(returns_by_role_data['clerk']),
                'items': returns_by_role_data['clerk'][:5],
            },
            'manager': {
                'count': len(returns_by_role_data['manager']),
                'items': returns_by_role_data['manager'][:5],
            },
            'purchaser': {
                'count': len(returns_by_role_data['purchaser']),
                'items': returns_by_role_data['purchaser'][:5],
            },
        },
        'visits_by_role': {
            'clerk': {
                'count': len(visits_by_role_data['clerk']),
                'items': visits_by_role_data['clerk'][:5],
            },
            'manager': {
                'count': len(visits_by_role_data['manager']),
                'items': visits_by_role_data['manager'][:5],
            },
            'purchaser': {
                'count': len(visits_by_role_data['purchaser']),
                'items': visits_by_role_data['purchaser'][:5],
            },
        },
        'statistics': stats,
    }


class ReturnExchangeCreateSchema(BaseModel):
    member_id: int
    type: str
    original_product_id: int
    original_batch_id: Optional[int] = None
    quantity: int = 1
    reason_category: str
    reason_detail: str
    exchange_product_id: Optional[int] = None
    assigned_to_id: Optional[int] = None


class ReturnExchangeUpdateSchema(BaseModel):
    status: Optional[str] = None
    assigned_to_id: Optional[int] = None
    stuck_reason: Optional[str] = None
    clerk_notes: Optional[str] = None
    manager_notes: Optional[str] = None
    purchaser_notes: Optional[str] = None
    amount_refunded: Optional[float] = None


@router.get("/returns")
def list_returns(
    request,
    status: Optional[str] = None,
    assigned_to: Optional[int] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
):
    """查询退换货列表"""
    queryset = ReturnExchange.objects.select_related(
        'member', 'assigned_to', 'assigned_to__user',
        'original_product', 'exchange_product'
    ).all()

    if status:
        queryset = queryset.filter(status=status)
    if assigned_to:
        queryset = queryset.filter(assigned_to_id=assigned_to)
    if date_from:
        queryset = queryset.filter(created_at__date__gte=date_from)
    if date_to:
        queryset = queryset.filter(created_at__date__lte=date_to)

    returns = []
    for ret in queryset[:50]:
        stuck_info = ret.get_stuck_info()
        returns.append({
            'id': ret.id,
            'return_number': ret.return_number,
            'member': {
                'id': ret.member.id,
                'name': ret.member.name,
                'phone': ret.member.phone,
            },
            'type': ret.get_type_display(),
            'status': ret.get_status_display(),
            'reason_category': ret.get_reason_category_display(),
            'reason_detail': ret.reason_detail,
            'original_product': {
                'id': ret.original_product.id,
                'product_code': ret.original_product.product_code,
                'name': ret.original_product.name,
                'price': float(ret.original_product.price),
            } if ret.original_product else None,
            'exchange_product': {
                'id': ret.exchange_product.id,
                'product_code': ret.exchange_product.product_code,
                'name': ret.exchange_product.name,
                'price': float(ret.exchange_product.price),
            } if ret.exchange_product else None,
            'quantity': ret.quantity,
            'assigned_to': {
                'id': ret.assigned_to.id,
                'name': ret.assigned_to.user.get_full_name(),
                'role': ret.assigned_to.get_role_display(),
            } if ret.assigned_to else None,
            'stuck_info': stuck_info,
            'created_at': ret.created_at,
        })

    return {'returns': returns, 'count': queryset.count()}


@router.get("/returns/{return_id}")
def get_return_detail(request, return_id: int):
    """获取退换货详情"""
    try:
        ret = ReturnExchange.objects.select_related(
            'member', 'assigned_to', 'assigned_to__user',
            'original_product', 'original_batch', 'exchange_product',
            'original_batch__product'
        ).get(id=return_id)

        return {
            'id': ret.id,
            'return_number': ret.return_number,
            'member': {
                'id': ret.member.id,
                'member_id': ret.member.member_id,
                'name': ret.member.name,
                'phone': ret.member.phone,
                'member_level': ret.member.get_member_level_display(),
            },
            'type': ret.type,
            'type_display': ret.get_type_display(),
            'status': ret.status,
            'status_display': ret.get_status_display(),
            'reason_category': ret.reason_category,
            'reason_category_display': ret.get_reason_category_display(),
            'reason_detail': ret.reason_detail,
            'original_product': {
                'id': ret.original_product.id,
                'product_code': ret.original_product.product_code,
                'name': ret.original_product.name,
                'price': float(ret.original_product.price),
            } if ret.original_product else None,
            'original_batch': {
                'id': ret.original_batch.id,
                'batch_number': ret.original_batch.batch_number,
                'production_date': ret.original_batch.production_date,
                'expiry_date': ret.original_batch.expiry_date,
            } if ret.original_batch else None,
            'exchange_product': {
                'id': ret.exchange_product.id,
                'product_code': ret.exchange_product.product_code,
                'name': ret.exchange_product.name,
                'price': float(ret.exchange_product.price),
            } if ret.exchange_product else None,
            'quantity': ret.quantity,
            'assigned_to': {
                'id': ret.assigned_to.id,
                'employee_id': ret.assigned_to.employee_id,
                'name': ret.assigned_to.user.get_full_name(),
                'role': ret.assigned_to.role,
                'role_display': ret.assigned_to.get_role_display(),
            } if ret.assigned_to else None,
            'current_handler_role': ret.get_current_expected_handler(),
            'stuck_reason': ret.stuck_reason,
            'stuck_info': ret.get_stuck_info(),
            'amount_refunded': float(ret.amount_refunded),
            'clerk_notes': ret.clerk_notes,
            'manager_notes': ret.manager_notes,
            'purchaser_notes': ret.purchaser_notes,
            'attachments': [
                {
                    'id': att.id,
                    'name': att.name,
                    'file_type': att.file_type,
                    'file_size': att.file_size,
                }
                for att in ret.attachments.all()
            ],
            'created_at': ret.created_at,
            'updated_at': ret.updated_at,
            'completed_at': ret.completed_at,
        }
    except ReturnExchange.DoesNotExist:
        return {'error': '退换货记录不存在'}


@router.post("/returns")
def create_return(request, data: ReturnExchangeCreateSchema):
    """创建退换货记录"""
    try:
        member = Member.objects.get(id=data.member_id)
        product = Product.objects.get(id=data.original_product_id)

        batch = None
        if data.original_batch_id:
            batch = ProductBatch.objects.get(id=data.original_batch_id)

        exchange_product = None
        if data.exchange_product_id:
            exchange_product = Product.objects.get(id=data.exchange_product_id)

        assigned_to = None
        if data.assigned_to_id:
            assigned_to = Employee.objects.get(id=data.assigned_to_id)

        today = timezone.now().date()
        count = ReturnExchange.objects.filter(
            created_at__date=today
        ).count() + 1

        return_number = f"RE{today.strftime('%Y%m%d')}{count:04d}"

        ret = ReturnExchange.objects.create(
            return_number=return_number,
            member=member,
            type=data.type,
            original_product=product,
            original_batch=batch,
            quantity=data.quantity,
            reason_category=data.reason_category,
            reason_detail=data.reason_detail,
            exchange_product=exchange_product,
            assigned_to=assigned_to,
            current_handler_role='clerk',
        )

        AuditLog.objects.create(
            action='create',
            entity_type='return_exchange',
            entity_id=str(ret.id),
            entity_name=return_number,
            user=assigned_to,
            description=f'创建退换货单 {return_number}',
        )

        return {
            'id': ret.id,
            'return_number': ret.return_number,
            'status': 'success',
        }
    except Exception as e:
        return {'error': str(e)}


@router.put("/returns/{return_id}")
def update_return(request, return_id: int, data: ReturnExchangeUpdateSchema):
    """更新退换货记录"""
    try:
        ret = ReturnExchange.objects.get(id=return_id)
        old_status = ret.status
        old_values = {'status': old_status}

        if data.status:
            ret.status = data.status
            if data.status == 'completed':
                ret.completed_at = timezone.now()

        if data.assigned_to_id:
            ret.assigned_to_id = data.assigned_to_id

        if data.stuck_reason is not None:
            ret.stuck_reason = data.stuck_reason

        if data.clerk_notes is not None:
            ret.clerk_notes = data.clerk_notes
        if data.manager_notes is not None:
            ret.manager_notes = data.manager_notes
        if data.purchaser_notes is not None:
            ret.purchaser_notes = data.purchaser_notes

        if data.amount_refunded is not None:
            ret.amount_refunded = data.amount_refunded

        ret.save()

        new_values = {
            'status': ret.status,
            'stuck_reason': ret.stuck_reason,
        }

        user = None
        if hasattr(request, 'user') and request.user.is_authenticated:
            try:
                user = request.user.employee_profile
            except Employee.DoesNotExist:
                pass

        AuditLog.objects.create(
            action='status_change' if data.status else 'update',
            entity_type='return_exchange',
            entity_id=str(ret.id),
            entity_name=ret.return_number,
            user=user,
            old_value=old_values,
            new_value=new_values,
            description=f'更新退换货单 {ret.return_number}',
        )

        return {
            'id': ret.id,
            'return_number': ret.return_number,
            'status': 'success',
        }
    except ReturnExchange.DoesNotExist:
        return {'error': '退换货记录不存在'}
    except Exception as e:
        return {'error': str(e)}


class VisitRecordCreateSchema(BaseModel):
    member_id: int
    type: str
    purpose: str
    priority: str = 'normal'
    assigned_to_id: Optional[int] = None
    scheduled_date: str
    scheduled_time: Optional[str] = None
    related_return_id: Optional[int] = None


class VisitRecordUpdateSchema(BaseModel):
    status: Optional[str] = None
    assigned_to_id: Optional[int] = None
    content: Optional[str] = None
    result: Optional[str] = None
    satisfaction_score: Optional[int] = None
    stuck_reason: Optional[str] = None


@router.get("/visits")
def list_visits(
    request,
    status: Optional[str] = None,
    assigned_to: Optional[int] = None,
    priority: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
):
    """查询客户回访列表"""
    queryset = VisitRecord.objects.select_related(
        'member', 'assigned_to', 'assigned_to__user'
    ).all()

    if status:
        queryset = queryset.filter(status=status)
    if assigned_to:
        queryset = queryset.filter(assigned_to_id=assigned_to)
    if priority:
        queryset = queryset.filter(priority=priority)
    if date_from:
        queryset = queryset.filter(scheduled_date__gte=date_from)
    if date_to:
        queryset = queryset.filter(scheduled_date__lte=date_to)

    visits = []
    for visit in queryset[:50]:
        stuck_info = visit.get_stuck_info()
        visits.append({
            'id': visit.id,
            'visit_number': visit.visit_number,
            'member': {
                'id': visit.member.id,
                'name': visit.member.name,
                'phone': visit.member.phone,
            },
            'type': visit.get_type_display(),
            'priority': visit.get_priority_display(),
            'status': visit.get_status_display(),
            'purpose': visit.purpose,
            'assigned_to': {
                'id': visit.assigned_to.id,
                'name': visit.assigned_to.user.get_full_name(),
                'role': visit.assigned_to.get_role_display(),
            } if visit.assigned_to else None,
            'scheduled_date': visit.scheduled_date,
            'scheduled_time': visit.scheduled_time,
            'stuck_info': stuck_info,
            'days_overdue': stuck_info.get('days_overdue') if stuck_info else None,
            'created_at': visit.created_at,
        })

    return {'visits': visits, 'count': queryset.count()}


@router.get("/visits/{visit_id}")
def get_visit_detail(request, visit_id: int):
    """获取客户回访详情"""
    try:
        visit = VisitRecord.objects.select_related(
            'member', 'assigned_to', 'assigned_to__user',
            'related_return'
        ).get(id=visit_id)

        return {
            'id': visit.id,
            'visit_number': visit.visit_number,
            'member': {
                'id': visit.member.id,
                'member_id': visit.member.member_id,
                'name': visit.member.name,
                'phone': visit.member.phone,
            },
            'type': visit.type,
            'type_display': visit.get_type_display(),
            'priority': visit.priority,
            'priority_display': visit.get_priority_display(),
            'status': visit.status,
            'status_display': visit.get_status_display(),
            'purpose': visit.purpose,
            'related_return': {
                'id': visit.related_return.id,
                'return_number': visit.related_return.return_number,
            } if visit.related_return else None,
            'assigned_to': {
                'id': visit.assigned_to.id,
                'employee_id': visit.assigned_to.employee_id,
                'name': visit.assigned_to.user.get_full_name(),
                'role': visit.assigned_to.role,
                'role_display': visit.assigned_to.get_role_display(),
            } if visit.assigned_to else None,
            'scheduled_date': visit.scheduled_date,
            'scheduled_time': visit.scheduled_time,
            'completed_date': visit.completed_date,
            'completed_time': visit.completed_time,
            'content': visit.content,
            'result': visit.result,
            'satisfaction_score': visit.satisfaction_score,
            'stuck_reason': visit.stuck_reason,
            'stuck_info': visit.get_stuck_info(),
            'attachments': [
                {
                    'id': att.id,
                    'name': att.name,
                    'file_type': att.file_type,
                    'file_size': att.file_size,
                }
                for att in visit.attachments.all()
            ],
            'created_at': visit.created_at,
            'updated_at': visit.updated_at,
        }
    except VisitRecord.DoesNotExist:
        return {'error': '回访记录不存在'}


@router.post("/visits")
def create_visit(request, data: VisitRecordCreateSchema):
    """创建客户回访记录"""
    try:
        member = Member.objects.get(id=data.member_id)

        scheduled_date = datetime.strptime(data.scheduled_date, '%Y-%m-%d').date()
        scheduled_time = None
        if data.scheduled_time:
            scheduled_time = datetime.strptime(data.scheduled_time, '%H:%M').time()

        assigned_to = None
        if data.assigned_to_id:
            assigned_to = Employee.objects.get(id=data.assigned_to_id)

        related_return = None
        if data.related_return_id:
            related_return = ReturnExchange.objects.get(id=data.related_return_id)

        today = timezone.now().date()
        count = VisitRecord.objects.filter(
            created_at__date=today
        ).count() + 1

        visit_number = f"VIS{today.strftime('%Y%m%d')}{count:04d}"

        visit = VisitRecord.objects.create(
            visit_number=visit_number,
            member=member,
            type=data.type,
            purpose=data.purpose,
            priority=data.priority,
            assigned_to=assigned_to,
            scheduled_date=scheduled_date,
            scheduled_time=scheduled_time,
            related_return=related_return,
        )

        AuditLog.objects.create(
            action='create',
            entity_type='visit_record',
            entity_id=str(visit.id),
            entity_name=visit_number,
            user=assigned_to,
            description=f'创建客户回访 {visit_number}',
        )

        return {
            'id': visit.id,
            'visit_number': visit.visit_number,
            'status': 'success',
        }
    except Exception as e:
        return {'error': str(e)}


@router.put("/visits/{visit_id}")
def update_visit(request, visit_id: int, data: VisitRecordUpdateSchema):
    """更新客户回访记录"""
    try:
        visit = VisitRecord.objects.get(id=visit_id)
        old_status = visit.status

        if data.status:
            visit.status = data.status
            if data.status == 'completed':
                visit.completed_date = timezone.now().date()
                visit.completed_time = timezone.now().time()

        if data.assigned_to_id:
            visit.assigned_to_id = data.assigned_to_id

        if data.content is not None:
            visit.content = data.content
        if data.result is not None:
            visit.result = data.result
        if data.satisfaction_score is not None:
            visit.satisfaction_score = data.satisfaction_score
        if data.stuck_reason is not None:
            visit.stuck_reason = data.stuck_reason

        visit.save()

        user = None
        if hasattr(request, 'user') and request.user.is_authenticated:
            try:
                user = request.user.employee_profile
            except Employee.DoesNotExist:
                pass

        AuditLog.objects.create(
            action='status_change' if data.status else 'update',
            entity_type='visit_record',
            entity_id=str(visit.id),
            entity_name=visit.visit_number,
            user=user,
            old_value={'status': old_status},
            new_value={'status': visit.status},
            description=f'更新客户回访 {visit.visit_number}',
        )

        return {
            'id': visit.id,
            'visit_number': visit.visit_number,
            'status': 'success',
        }
    except VisitRecord.DoesNotExist:
        return {'error': '回访记录不存在'}
    except Exception as e:
        return {'error': str(e)}


@router.get("/audit-logs")
def list_audit_logs(
    request,
    entity_type: Optional[str] = None,
    entity_id: Optional[int] = None,
    action: Optional[str] = None,
    user_id: Optional[int] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    limit: int = Query(default=50, le=200),
):
    """查询审计日志"""
    queryset = AuditLog.objects.select_related(
        'user', 'user__user'
    ).all()

    if entity_type:
        queryset = queryset.filter(entity_type=entity_type)
    if entity_id:
        queryset = queryset.filter(entity_id=str(entity_id))
    if action:
        queryset = queryset.filter(action=action)
    if user_id:
        queryset = queryset.filter(user_id=user_id)
    if date_from:
        queryset = queryset.filter(created_at__date__gte=date_from)
    if date_to:
        queryset = queryset.filter(created_at__date__lte=date_to)

    logs = []
    for log in queryset[:limit]:
        logs.append({
            'id': log.id,
            'action': log.action,
            'action_display': log.get_action_display(),
            'entity_type': log.entity_type,
            'entity_type_display': log.get_entity_type_display(),
            'entity_id': log.entity_id,
            'entity_name': log.entity_name,
            'user': {
                'id': log.user.id,
                'name': log.user.user.get_full_name() if log.user else None,
                'role': log.user.get_role_display() if log.user else None,
            } if log.user else None,
            'old_value': log.old_value,
            'new_value': log.new_value,
            'description': log.description,
            'created_at': log.created_at,
        })

    return {'logs': logs, 'count': queryset.count()}


@router.get("/audit-logs/export")
def export_audit_logs(
    request,
    entity_type: Optional[str] = None,
    entity_id: Optional[int] = None,
    action: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
):
    """导出审计日志到Excel"""
    queryset = AuditLog.objects.select_related(
        'user', 'user__user'
    ).all()

    if entity_type:
        queryset = queryset.filter(entity_type=entity_type)
    if entity_id:
        queryset = queryset.filter(entity_id=str(entity_id))
    if action:
        queryset = queryset.filter(action=action)
    if date_from:
        queryset = queryset.filter(created_at__date__gte=date_from)
    if date_to:
        queryset = queryset.filter(created_at__date__lte=date_to)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '审计日志'

    headers = ['操作时间', '操作类型', '实体类型', '实体ID', '实体名称', '操作用户', '用户角色', '描述', '旧值', '新值']
    ws.append(headers)

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for log in queryset[:1000]:
        ws.append([
            log.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            log.get_action_display(),
            log.get_entity_type_display(),
            log.entity_id,
            log.entity_name,
            log.user.user.get_full_name() if log.user else '',
            log.user.get_role_display() if log.user else '',
            log.description,
            str(log.old_value) if log.old_value else '',
            str(log.new_value) if log.new_value else '',
        ])

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    user = None
    if hasattr(request, 'user') and request.user.is_authenticated:
        try:
            user = request.user.employee_profile
        except Employee.DoesNotExist:
            pass

    AuditLog.objects.create(
        action='export',
        entity_type='audit_logs',
        entity_id='batch',
        user=user,
        description=f'导出审计日志，数量: {queryset.count()}',
    )

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=audit_logs_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    return response


@router.get("/returns/export")
def export_returns(
    request,
    status: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
):
    """导出退换货记录到Excel"""
    queryset = ReturnExchange.objects.select_related(
        'member', 'assigned_to', 'assigned_to__user',
        'original_product'
    ).all()

    if status:
        queryset = queryset.filter(status=status)
    if date_from:
        queryset = queryset.filter(created_at__date__gte=date_from)
    if date_to:
        queryset = queryset.filter(created_at__date__lte=date_to)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '退换货记录'

    headers = ['退换货单号', '会员姓名', '会员电话', '类型', '状态', '原因分类',
               '原商品', '数量', '处理人', '当前处理角色', '退款金额', '创建时间', '完成时间']
    ws.append(headers)

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='70AD47', end_color='70AD47', fill_type='solid')
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for ret in queryset[:1000]:
        ws.append([
            ret.return_number,
            ret.member.name,
            ret.member.phone,
            ret.get_type_display(),
            ret.get_status_display(),
            ret.get_reason_category_display(),
            ret.original_product.name if ret.original_product else '',
            ret.quantity,
            ret.assigned_to.user.get_full_name() if ret.assigned_to else '',
            ret.get_current_expected_handler(),
            float(ret.amount_refunded),
            ret.created_at.strftime('%Y-%m-%d %H:%M'),
            ret.completed_at.strftime('%Y-%m-%d %H:%M') if ret.completed_at else '',
        ])

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    user = None
    if hasattr(request, 'user') and request.user.is_authenticated:
        try:
            user = request.user.employee_profile
        except Employee.DoesNotExist:
            pass

    AuditLog.objects.create(
        action='export',
        entity_type='return_exchange',
        entity_id='batch',
        user=user,
        description=f'导出退换货记录，数量: {queryset.count()}',
    )

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=returns_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    return response


@router.get("/visits/export")
def export_visits(
    request,
    status: Optional[str] = None,
    priority: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
):
    """导出客户回访记录到Excel"""
    queryset = VisitRecord.objects.select_related(
        'member', 'assigned_to', 'assigned_to__user'
    ).all()

    if status:
        queryset = queryset.filter(status=status)
    if priority:
        queryset = queryset.filter(priority=priority)
    if date_from:
        queryset = queryset.filter(scheduled_date__gte=date_from)
    if date_to:
        queryset = queryset.filter(scheduled_date__lte=date_to)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '客户回访记录'

    headers = ['回访编号', '会员姓名', '会员电话', '回访方式', '优先级', '状态',
               '计划日期', '计划时间', '回访人', '满意度', '创建时间', '完成时间']
    ws.append(headers)

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='ED7D31', end_color='ED7D31', fill_type='solid')
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for visit in queryset[:1000]:
        ws.append([
            visit.visit_number,
            visit.member.name,
            visit.member.phone,
            visit.get_type_display(),
            visit.get_priority_display(),
            visit.get_status_display(),
            str(visit.scheduled_date),
            str(visit.scheduled_time) if visit.scheduled_time else '',
            visit.assigned_to.user.get_full_name() if visit.assigned_to else '',
            visit.satisfaction_score if visit.satisfaction_score else '',
            visit.created_at.strftime('%Y-%m-%d %H:%M'),
            visit.completed_date.strftime('%Y-%m-%d') if visit.completed_date else '',
        ])

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    user = None
    if hasattr(request, 'user') and request.user.is_authenticated:
        try:
            user = request.user.employee_profile
        except Employee.DoesNotExist:
            pass

    AuditLog.objects.create(
        action='export',
        entity_type='visit_record',
        entity_id='batch',
        user=user,
        description=f'导出客户回访记录，数量: {queryset.count()}',
    )

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=visits_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    return response


@router.get("/notifications")
def list_notifications(request, is_read: Optional[bool] = None, limit: int = Query(default=20, le=100)):
    """获取消息通知列表"""
    try:
        employee = request.user.employee_profile
        queryset = Notification.objects.filter(recipient=employee)

        if is_read is not None:
            queryset = queryset.filter(is_read=is_read)

        notifications = []
        for notif in queryset[:limit]:
            notifications.append({
                'id': notif.id,
                'type': notif.type,
                'type_display': notif.get_type_display(),
                'title': notif.title,
                'content': notif.content,
                'sender': {
                    'name': notif.sender.user.get_full_name() if notif.sender else None,
                } if notif.sender else None,
                'is_read': notif.is_read,
                'read_at': notif.read_at,
                'related_entity_type': notif.related_entity_type,
                'related_entity_id': notif.related_entity_id,
                'created_at': notif.created_at,
            })

        return {
            'notifications': notifications,
            'unread_count': queryset.filter(is_read=False).count(),
        }
    except Employee.DoesNotExist:
        return {'error': '用户不存在', 'notifications': [], 'unread_count': 0}


@router.put("/notifications/{notification_id}/read")
def mark_notification_read(request, notification_id: int):
    """标记通知为已读"""
    try:
        employee = request.user.employee_profile
        notif = Notification.objects.get(id=notification_id, recipient=employee)
        notif.is_read = True
        notif.read_at = timezone.now()
        notif.save()
        return {'status': 'success'}
    except Notification.DoesNotExist:
        return {'error': '通知不存在'}


@router.post("/notifications")
def create_notification(request, notification_data: dict):
    """创建消息通知"""
    try:
        notification = Notification.objects.create(
            type=notification_data.get('type', 'system'),
            channel=notification_data.get('channel', 'internal'),
            recipient_id=notification_data['recipient_id'],
            sender_id=notification_data.get('sender_id'),
            title=notification_data['title'],
            content=notification_data['content'],
            related_entity_type=notification_data.get('related_entity_type', ''),
            related_entity_id=notification_data.get('related_entity_id', ''),
        )
        return {
            'id': notification.id,
            'status': 'success',
        }
    except Exception as e:
        return {'error': str(e)}


@router.get("/employees")
def list_employees(request, role: Optional[str] = None):
    """获取员工列表"""
    queryset = Employee.objects.select_related('user').filter(is_active=True)

    if role:
        queryset = queryset.filter(role=role)

    employees = []
    for emp in queryset:
        employees.append({
            'id': emp.id,
            'employee_id': emp.employee_id,
            'name': emp.user.get_full_name() or emp.user.username,
            'role': emp.role,
            'role_display': emp.get_role_display(),
            'phone': emp.phone,
        })

    return {'employees': employees}


@router.get("/members")
def list_members(
    request,
    search: Optional[str] = None,
    member_level: Optional[str] = None,
):
    """获取会员列表"""
    queryset = Member.objects.all()

    if search:
        queryset = queryset.filter(
            Q(name__icontains=search) | Q(phone__icontains=search) | Q(member_id__icontains=search)
        )
    if member_level:
        queryset = queryset.filter(member_level=member_level)

    members = []
    for member in queryset[:50]:
        members.append({
            'id': member.id,
            'member_id': member.member_id,
            'name': member.name,
            'phone': member.phone,
            'gender': member.gender,
            'member_level': member.get_member_level_display(),
            'total_points': member.total_points,
        })

    return {'members': members, 'count': queryset.count()}